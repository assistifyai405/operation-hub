"""Proposal & Contract exporters — branded professional PDF (reportlab) and DOCX (python-docx)."""
import io
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, HRFlowable,
    Image as RLImage, PageBreak, Table, TableStyle,
)
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from proposal_config import PROPOSAL_SECTIONS

ASSISTIFY_GREEN = "#16A34A"


def _as_list(v):
    if isinstance(v, list):
        return v
    return [v] if v else []


def _hex_to_rgb(h):
    h = (h or ASSISTIFY_GREEN).lstrip("#")
    if len(h) != 6:
        h = ASSISTIFY_GREEN.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _default_brand(brand):
    b = brand or {}
    return {
        "company_name": b.get("company_name", "Assistify OS"),
        "logo_bytes": b.get("logo_bytes"),
        "primary": b.get("primary") or ASSISTIFY_GREEN,
        "secondary": b.get("secondary") or "#22D3EE",
        "address": b.get("address", ""), "website": b.get("website", ""),
        "email": b.get("email", ""), "phone": b.get("phone", ""),
        "vat": b.get("vat", ""), "kvk": b.get("kvk", ""), "footer": b.get("footer", ""),
    }


def _logo_flowable(brand, max_w=1.7 * inch, max_h=0.9 * inch):
    if not brand.get("logo_bytes"):
        return None
    try:
        reader = ImageReader(io.BytesIO(brand["logo_bytes"]))
        iw, ih = reader.getSize()
        ratio = min(max_w / iw, max_h / ih)
        return RLImage(io.BytesIO(brand["logo_bytes"]), width=iw * ratio, height=ih * ratio)
    except Exception:
        return None


def _company_lines(brand):
    parts = []
    for key in ("address", "email", "phone", "website"):
        if brand.get(key):
            parts.append(str(brand[key]))
    ids = []
    if brand.get("vat"):
        ids.append(f"VAT: {brand['vat']}")
    if brand.get("kvk"):
        ids.append(f"KVK: {brand['kvk']}")
    if ids:
        parts.append(" · ".join(ids))
    return parts


def _make_footer(brand):
    accent = HexColor(brand["primary"])
    footer_text = brand.get("footer") or brand.get("company_name", "")

    def _draw(canvas, doc):
        canvas.saveState()
        w, h = LETTER
        canvas.setStrokeColor(accent)
        canvas.setLineWidth(0.8)
        canvas.line(0.9 * inch, 0.62 * inch, w - 0.9 * inch, 0.62 * inch)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(HexColor("#71717A"))
        if footer_text:
            canvas.drawString(0.9 * inch, 0.45 * inch, footer_text[:110])
        canvas.drawRightString(w - 0.9 * inch, 0.45 * inch, f"Page {doc.page}")
        canvas.restoreState()

    return _draw


def build_pdf(proposal: dict, sections=PROPOSAL_SECTIONS, doc_type: str = "Proposal", brand=None, signature=False) -> bytes:
    brand = _default_brand(brand)
    accent = HexColor(brand["primary"])
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, topMargin=0.9 * inch, bottomMargin=0.9 * inch,
                            leftMargin=0.9 * inch, rightMargin=0.9 * inch)
    ss = getSampleStyleSheet()
    cover_title = ParagraphStyle("ct", parent=ss["Title"], fontSize=32, textColor=HexColor("#18181B"), spaceAfter=6, leading=38)
    cover_sub = ParagraphStyle("cs", parent=ss["Normal"], fontSize=12, textColor=accent, spaceAfter=4)
    company = ParagraphStyle("co", parent=ss["Normal"], fontSize=15, textColor=HexColor("#18181B"), spaceAfter=2, leading=19)
    small = ParagraphStyle("sm", parent=ss["Normal"], fontSize=9.5, textColor=HexColor("#71717A"), leading=14)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=13, textColor=accent, spaceBefore=16, spaceAfter=6)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=10.5, leading=15, textColor=HexColor("#27272A"), alignment=TA_LEFT)

    content = proposal.get("content", {})
    story = []

    # ---- Cover page ----
    logo = _logo_flowable(brand)
    if logo:
        story += [logo, Spacer(1, 26)]
    else:
        story += [Paragraph(brand["company_name"], company), Spacer(1, 26)]
    story += [Spacer(1, 60)]
    story += [Paragraph(proposal.get("title", doc_type), cover_title)]
    story += [Paragraph(f"{doc_type} · v{proposal.get('version', 1)} · {proposal.get('status', 'Draft')}", cover_sub)]
    story += [HRFlowable(width="40%", color=accent, thickness=2, spaceBefore=10, spaceAfter=20, hAlign="LEFT")]
    if logo:
        story += [Paragraph(f"<b>{brand['company_name']}</b>", company)]
    for line in _company_lines(brand):
        story.append(Paragraph(line, small))
    story += [PageBreak()]

    # ---- Body ----
    for sec in sections:
        val = content.get(sec["key"])
        if not val:
            continue
        story.append(Paragraph(sec["label"], h2))
        if sec["type"] == "list":
            items = [ListItem(Paragraph(str(i), body), leftIndent=8) for i in _as_list(val)]
            if items:
                story.append(ListFlowable(items, bulletType="bullet", start="•"))
        else:
            story.append(Paragraph(str(val).replace("\n", "<br/>"), body))
        story.append(Spacer(1, 4))

    # ---- Signature page (contracts) ----
    if signature:
        story += [PageBreak(), Paragraph("Signatures", h2), Spacer(1, 10),
                  Paragraph("By signing below, both parties agree to the terms set out in this agreement.", body), Spacer(1, 40)]
        sig = [
            [Paragraph("_______________________________", body), Paragraph("_______________________________", body)],
            [Paragraph(f"<b>{brand['company_name']}</b> (Provider)", small), Paragraph("Client", small)],
            [Spacer(1, 24), Spacer(1, 24)],
            [Paragraph("Name: _______________________", small), Paragraph("Name: _______________________", small)],
            [Paragraph("Date: _______________________", small), Paragraph("Date: _______________________", small)],
        ]
        st = Table(sig, colWidths=[3.3 * inch, 3.3 * inch])
        st.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6)]))
        story.append(st)

    footer = _make_footer(brand)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def build_docx(proposal: dict, sections=PROPOSAL_SECTIONS, doc_type: str = "Proposal", brand=None, signature=False) -> bytes:
    brand = _default_brand(brand)
    accent = _hex_to_rgb(brand["primary"])
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    if brand.get("logo_bytes"):
        try:
            d.add_picture(io.BytesIO(brand["logo_bytes"]), width=Inches(1.7))
        except Exception:
            pass
    cp = d.add_paragraph(); cr = cp.add_run(brand["company_name"]); cr.bold = True; cr.font.size = Pt(13)
    for line in _company_lines(brand):
        lp = d.add_paragraph(); lr = lp.add_run(line); lr.font.size = Pt(9); lr.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    content = proposal.get("content", {})
    t = d.add_paragraph()
    run = t.add_run(proposal.get("title", doc_type))
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x18, 0x18, 0x1B)

    meta = d.add_paragraph()
    mr = meta.add_run(f"{doc_type} v{proposal.get('version', 1)}  ·  Status: {proposal.get('status', 'Draft')}")
    mr.font.size = Pt(10); mr.font.color.rgb = accent

    for sec in sections:
        val = content.get(sec["key"])
        if not val:
            continue
        h = d.add_heading(level=2)
        hr = h.add_run(sec["label"]); hr.font.color.rgb = accent
        if sec["type"] == "list":
            for item in _as_list(val):
                d.add_paragraph(str(item), style="List Bullet")
        else:
            d.add_paragraph(str(val))

    if signature:
        d.add_page_break()
        sh = d.add_heading(level=2); shr = sh.add_run("Signatures"); shr.font.color.rgb = accent
        d.add_paragraph("By signing below, both parties agree to the terms set out in this agreement.")
        d.add_paragraph("\n")
        d.add_paragraph(f"{brand['company_name']} (Provider): _______________________    Date: __________")
        d.add_paragraph("Client: _______________________    Date: __________")

    if brand.get("footer"):
        section = d.sections[0]
        fp = section.footer.paragraphs[0]
        fr = fp.add_run(brand["footer"]); fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
