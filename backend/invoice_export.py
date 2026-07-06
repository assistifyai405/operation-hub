"""Invoice exporters — branded professional PDF (reportlab) and DOCX (python-docx)."""
import io
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, Image as RLImage
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

VIOLET = "#7C3AED"


def _money(v):
    try:
        return f"{float(v):,.2f}"
    except Exception:
        return "0.00"


def _rows(invoice):
    return invoice.get("line_items", []) or []


def _hex_to_rgb(h):
    h = (h or "#7C3AED").lstrip("#")
    if len(h) != 6:
        h = "7C3AED"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _default_brand(brand):
    b = brand or {}
    return {
        "company_name": b.get("company_name", "Assistify OS"),
        "logo_bytes": b.get("logo_bytes"),
        "primary": b.get("primary") or VIOLET,
        "address": b.get("address", ""), "website": b.get("website", ""),
        "email": b.get("email", ""), "phone": b.get("phone", ""),
        "vat": b.get("vat", ""), "kvk": b.get("kvk", ""), "footer": b.get("footer", ""),
    }


def _company_lines(brand):
    parts = []
    for key in ("address", "email", "phone", "website"):
        if brand.get(key):
            parts.append(str(brand[key]))
    ids = []
    if brand.get("vat"):
        ids.append(f"VAT: {brand['vat']}")
    if brand.get("kvk"):
        ids.append(f"KVK: {brand['kvk']}")
    if ids:
        parts.append(" · ".join(ids))
    return parts


def _make_footer(brand):
    accent = HexColor(brand["primary"])
    footer_text = brand.get("footer") or brand.get("company_name", "")

    def _draw(canvas, doc):
        canvas.saveState()
        w, h = LETTER
        canvas.setStrokeColor(accent)
        canvas.setLineWidth(0.8)
        canvas.line(0.8 * inch, 0.55 * inch, w - 0.8 * inch, 0.55 * inch)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(HexColor("#71717A"))
        if footer_text:
            canvas.drawString(0.8 * inch, 0.38 * inch, footer_text[:110])
        canvas.drawRightString(w - 0.8 * inch, 0.38 * inch, f"Page {doc.page}")
        canvas.restoreState()

    return _draw


def build_invoice_pdf(invoice: dict, brand=None) -> bytes:
    brand = _default_brand(brand)
    accent = HexColor(brand["primary"])
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, topMargin=0.8 * inch, bottomMargin=0.85 * inch,
                            leftMargin=0.8 * inch, rightMargin=0.8 * inch)
    ss = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=ss["Title"], fontSize=28, textColor=accent, spaceAfter=0, alignment=2)
    small = ParagraphStyle("s", parent=ss["Normal"], fontSize=9.5, textColor=HexColor("#52525B"), leading=13)
    smallr = ParagraphStyle("sr", parent=small, alignment=2)
    label = ParagraphStyle("l", parent=ss["Normal"], fontSize=9, textColor=HexColor("#A1A1AA"), spaceAfter=1)
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=10, textColor=HexColor("#27272A"), leading=14)
    comp = ParagraphStyle("cp", parent=ss["Normal"], fontSize=11, textColor=HexColor("#18181B"), leading=15)
    c = invoice.get("content", {})

    # ---- Branded header: logo/company left, INVOICE right ----
    try:
        reader = ImageReader(io.BytesIO(brand["logo_bytes"])) if brand.get("logo_bytes") else None
    except Exception:
        reader = None
    if reader:
        iw, ih = reader.getSize()
        ratio = min((1.6 * inch) / iw, (0.7 * inch) / ih)
        left = RLImage(io.BytesIO(brand["logo_bytes"]), width=iw * ratio, height=ih * ratio)
    else:
        left = Paragraph(f"<b>{brand['company_name']}</b>", comp)
    right = Paragraph("INVOICE", title)
    header = Table([[left, right]], colWidths=[3.6 * inch, 3.2 * inch])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story = [header]
    comp_block = "<br/>".join([f"<b>{brand['company_name']}</b>"] + _company_lines(brand)) if reader else "<br/>".join(_company_lines(brand))
    story += [Table([[Paragraph(comp_block, small),
                      Paragraph(f"{invoice.get('invoice_number', '')}<br/>Status: {invoice.get('status', 'Draft')}", smallr)]],
                    colWidths=[3.6 * inch, 3.2 * inch])]
    story += [Spacer(1, 10), HRFlowable(width="100%", color=accent, thickness=1.2), Spacer(1, 12)]

    meta = [
        [Paragraph("BILL TO", label), Paragraph("DETAILS", label)],
        [Paragraph(f"<b>{c.get('company') or c.get('client_name') or '—'}</b><br/>{(c.get('billing_address') or '').replace(chr(10), '<br/>')}", body),
         Paragraph(f"Issue Date: {c.get('issue_date') or '—'}<br/>Due Date: {c.get('due_date') or '—'}<br/>Project: {c.get('project_name') or '—'}", body)],
    ]
    mt = Table(meta, colWidths=[3.4 * inch, 3.4 * inch])
    mt.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, 0), 4)]))
    story += [mt, Spacer(1, 8)]
    if c.get("description"):
        story += [Paragraph(c["description"], body), Spacer(1, 10)]

    data = [["Description", "Qty", "Unit Price", "Amount"]]
    for li in _rows(invoice):
        data.append([Paragraph(str(li.get("description", "")), body), str(li.get("quantity", 0)),
                     _money(li.get("unit_price", 0)), _money(li.get("amount", 0))])
    tbl = Table(data, colWidths=[3.6 * inch, 0.7 * inch, 1.25 * inch, 1.25 * inch])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), accent),
        ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#FFFFFF")),
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#FFFFFF"), HexColor("#F4F4F5")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.5, HexColor("#E4E4E7")),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story += [tbl, Spacer(1, 10)]

    totals = [["Subtotal", _money(invoice.get("subtotal", 0))],
              [f"VAT ({c.get('vat_rate', 0)}%)", _money(invoice.get("vat", 0))],
              ["Total", _money(invoice.get("total", 0))]]
    tt = Table(totals, colWidths=[1.5 * inch, 1.25 * inch], hAlign="RIGHT")
    tt.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "RIGHT"), ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LINEABOVE", (0, -1), (-1, -1), 1, accent),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [tt, Spacer(1, 16)]

    if c.get("payment_terms") or c.get("bank_details"):
        story += [Paragraph("Payment Instructions", label)]
        for key in ("payment_terms", "bank_details"):
            if c.get(key):
                story += [Paragraph(f"<b>{key.replace('_', ' ').title()}:</b> {str(c[key]).replace(chr(10), '<br/>')}", body)]
        story += [Spacer(1, 6)]
    if c.get("notes"):
        story += [Paragraph("Notes", label), Paragraph(str(c["notes"]).replace("\n", "<br/>"), body)]

    footer = _make_footer(brand)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def build_invoice_docx(invoice: dict, brand=None) -> bytes:
    brand = _default_brand(brand)
    accent = _hex_to_rgb(brand["primary"])
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    c = invoice.get("content", {})

    if brand.get("logo_bytes"):
        try:
            d.add_picture(io.BytesIO(brand["logo_bytes"]), width=Inches(1.6))
        except Exception:
            pass
    cp = d.add_paragraph(); cr = cp.add_run(brand["company_name"]); cr.bold = True; cr.font.size = Pt(12)
    for line in _company_lines(brand):
        lp = d.add_paragraph(); lr = lp.add_run(line); lr.font.size = Pt(9); lr.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    h = d.add_paragraph()
    r = h.add_run("INVOICE")
    r.bold = True; r.font.size = Pt(26); r.font.color.rgb = accent
    m = d.add_paragraph(); mr = m.add_run(f"{invoice.get('invoice_number', '')}  ·  Status: {invoice.get('status', 'Draft')}")
    mr.font.size = Pt(10); mr.font.color.rgb = RGBColor(0x52, 0x52, 0x5B)

    d.add_paragraph(f"Bill To: {c.get('company') or c.get('client_name') or '—'}")
    if c.get("billing_address"):
        d.add_paragraph(str(c["billing_address"]))
    d.add_paragraph(f"Issue Date: {c.get('issue_date') or '—'}    Due Date: {c.get('due_date') or '—'}    Project: {c.get('project_name') or '—'}")
    if c.get("description"):
        d.add_paragraph(str(c["description"]))

    table = d.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 4"
    hdr = table.rows[0].cells
    for i, t in enumerate(["Description", "Qty", "Unit Price", "Amount"]):
        hdr[i].text = t
    for li in _rows(invoice):
        row = table.add_row().cells
        row[0].text = str(li.get("description", ""))
        row[1].text = str(li.get("quantity", 0))
        row[2].text = _money(li.get("unit_price", 0))
        row[3].text = _money(li.get("amount", 0))

    d.add_paragraph()
    for lbl, val in [("Subtotal", invoice.get("subtotal", 0)), (f"VAT ({c.get('vat_rate', 0)}%)", invoice.get("vat", 0)), ("Total", invoice.get("total", 0))]:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(f"{lbl}: {_money(val)}")
        if lbl == "Total":
            run.bold = True; run.font.color.rgb = accent

    if c.get("payment_terms") or c.get("bank_details"):
        hp = d.add_paragraph(); hr = hp.add_run("Payment Instructions"); hr.bold = True
        for key in ("payment_terms", "bank_details"):
            if c.get(key):
                d.add_paragraph(f"{key.replace('_', ' ').title()}: {c[key]}")
    if c.get("notes"):
        hp = d.add_paragraph(); hr = hp.add_run("Notes"); hr.bold = True
        d.add_paragraph(str(c["notes"]))

    if brand.get("footer"):
        fp = d.sections[0].footer.paragraphs[0]
        fr = fp.add_run(brand["footer"]); fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
